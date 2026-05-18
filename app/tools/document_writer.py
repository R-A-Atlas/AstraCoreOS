from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re

from docx import Document


@dataclass(frozen=True)
class DocumentArtifact:
    title: str
    path: Path
    kind: str
    summary: str


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", text.strip().lower()).strip("-")
    return slug[:80] or "document"


def build_mobile_detailing_plan() -> str:
    return """# Mobile Detailing Business Plan

## Executive Summary
A lean mobile detailing business can start with low overhead, local lead generation, and service packages that are easy to fulfill repeatedly. The first target should be residential customers, busy professionals, rideshare drivers, and small fleets.

## Offer
- Basic wash and interior reset
- Premium interior and exterior detail
- Ceramic spray protection add-on
- Monthly maintenance plan
- Fleet cleaning for small businesses

## Target Customer
The best early customers are people who value convenience more than the lowest price: professionals, parents, real estate agents, car enthusiasts, and drivers who earn income from their vehicle.

## Startup Needs
- Pressure washer or rinseless wash setup
- Vacuum, extractor, microfiber towels, brushes, chemicals
- Portable power/water plan depending on local rules
- Booking page, Google Business Profile, before/after photo system
- Liability insurance before taking fleet or high-value jobs

## Pricing Model
Start with three clear packages:
- Essential: entry-level maintenance
- Signature: full interior/exterior detail
- Elite: deep clean plus protection

Avoid too many custom prices early. Simple packages make sales and scheduling easier.

## Marketing Plan
Launch with neighborhood Facebook groups, Google Business Profile, short-form before/after videos, referral cards, and partnerships with apartments, gyms, offices, and small dealerships.

## Operations
Use a repeatable checklist for each job. Track time per package, supplies used, customer notes, photos, and follow-up date. The first operational goal is consistent quality, not maximum volume.

## 30-Day Plan
Week 1: buy core supplies, create packages, set up Google Business Profile.
Week 2: complete five discounted portfolio jobs and collect photos/reviews.
Week 3: launch local outreach and referral offer.
Week 4: convert best customers into monthly maintenance plans.

## Risks
Main risks are underpricing, inconsistent quality, weather delays, weak scheduling, and not collecting reviews. Protect margins by measuring time per job and raising prices quickly when demand proves out.

## Bottom Line
This business is attractive because it can start small, generate cash quickly, and compound through reviews, recurring customers, and local trust. The first milestone is ten paid jobs with documented before/after proof and at least five public reviews.
"""


def build_business_document_markdown(title: str, packet_data: dict) -> str:
    business = packet_data.get("business", "the business")
    document_type = str(packet_data.get("document_type", "document")).replace("_", " ")
    sections = packet_data.get("recommended_sections", [])
    assumptions = packet_data.get("operating_assumptions", [])
    quality_bar = packet_data.get("quality_bar", [])

    lines = [
        f"# {title}",
        "",
        "## Overview",
        f"This {document_type} is built for {business}. It is intentionally practical: clear structure, useful assumptions, and next steps a real operator can act on.",
        "",
    ]

    if assumptions:
        lines.extend(["## Working Assumptions"])
        lines.extend(f"- {item}" for item in assumptions)
        lines.append("")

    for section in sections:
        lines.extend([f"## {section}", _section_body(section, business, document_type), ""])

    if quality_bar:
        lines.extend(["## Quality Bar"])
        lines.extend(f"- {item}" for item in quality_bar)
        lines.append("")

    return "\n".join(lines).strip() + "\n"


def _section_body(section: str, business: str, document_type: str) -> str:
    templates = {
        "Executive Summary": f"{business.title()} should be positioned around a simple offer, repeatable delivery, and clear proof that the service solves a real customer problem.",
        "Offer": "Define 2-3 clear packages. Each package should have a fixed scope, expected delivery time, and obvious upgrade path.",
        "Target Customer": "Start with the customer segment that values reliability and convenience enough to pay without heavy negotiation.",
        "Startup Needs": "List required tools, software, compliance items, insurance, payment setup, and the minimum marketing assets needed to start selling.",
        "Pricing Model": "Use simple tiers first. Track time, materials, customer acquisition cost, and gross margin before adding complexity.",
        "Marketing Plan": "Prioritize channels that create visible proof: referrals, local search, before/after content, partnerships, and direct outreach.",
        "Operations": "Create a repeatable checklist, collect customer notes, measure cycle time, and keep a follow-up rhythm.",
        "30-Day Plan": "Week 1: setup. Week 2: first proof jobs. Week 3: outreach. Week 4: convert best leads into repeat business.",
        "Risks": "Main risks are unclear scope, weak pricing, inconsistent quality, no review system, and missing compliance or insurance requirements.",
        "Bottom Line": f"The first milestone is not scale. The first milestone is proving that {business} can produce consistent outcomes, margin, and repeat demand.",
        "Client Problem": "State the client's business problem in plain language before describing the solution.",
        "Recommended Solution": "Connect the proposed work directly to revenue, savings, speed, trust, or operational clarity.",
        "Scope of Work": "Define exactly what is included, what is excluded, and what requires a change order.",
        "Timeline": "Break delivery into short phases with review points and clear owner responsibilities.",
        "Investment": "Show the price, payment timing, and what the client receives for that investment.",
        "Success Metrics": "Define how both sides will know the work succeeded.",
        "Next Step": "Give one concrete next action: approve, schedule, provide assets, or pay deposit.",
        "Purpose": f"This SOP standardizes how {business} completes the work without relying on memory.",
        "Required Materials": "List tools, templates, access, supplies, safety items, and pre-work checks.",
        "Step-by-Step Procedure": "Write the process in order from intake to delivery, including handoff points.",
        "Quality Check": "Define what must be true before the work is considered complete.",
        "Common Failure Points": "Name the most likely mistakes and how to prevent them.",
        "Owner": "Assign the person responsible for keeping this process current.",
        "Snapshot": "Summarize current position, key numbers needed, and the immediate decision this report supports.",
        "Revenue Drivers": "Identify the few activities most likely to create revenue in the next 30-90 days.",
        "Cost Structure": "Separate fixed costs, variable costs, startup costs, and optional costs.",
        "Cash Needs": "Estimate what must be paid before revenue arrives and where cash could get tight.",
        "Next 30 Days": "Focus on actions that prove demand, reduce uncertainty, and create measurable progress.",
        "Context": "State what is being evaluated and why it matters now.",
        "Key Findings": "List the highest-signal findings first.",
        "Operational Impact": "Explain what changes in decisions, workload, costs, or risk.",
        "Recommended Next Steps": "Prioritize the smallest useful actions before larger commitments.",
    }
    return templates.get(section, f"Use this section to make the {document_type} specific to {business}.")


def write_docx(title: str, markdown: str, output_dir: Path, summary: str | None = None) -> DocumentArtifact:
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    path = output_dir / f"{slugify(title)}-{stamp}.docx"

    doc = Document()
    doc.add_heading(title, level=0)

    for raw in markdown.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)

    doc.save(path)
    return DocumentArtifact(
        title=title,
        path=path,
        kind="docx",
        summary=summary or "Created a Word document from the gathered context packet.",
    )
